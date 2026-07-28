"""
apps/products/importers_extended.py — Importers para Brand y Course.

Reutiliza los helpers de importers.py (normalize_header, parse_bool, parse_decimal).
"""
import csv
import io
import openpyxl
from django.utils.text import slugify

from .importers import normalize_header, parse_bool, parse_decimal
from .models    import Brand


# ═══════════════════════════════════════════════════════════════
# BRANDS
# ═══════════════════════════════════════════════════════════════

BRAND_COLUMN_MAP = {
    "nombre":     "name",
    "sitio_web":  "website",
    "website":    "website",
    "url":        "website",
    "logo_url":   "logo_url",
    "activo":     "is_active",
    "activa":     "is_active",
}
BRAND_REQUIRED = ["nombre"]


def _parse_brand_rows(rows: list[dict]) -> dict:
    created, updated, errors = [], [], []

    for i, row in enumerate(rows, start=2):
        name = str(row.get("name", "")).strip()
        if not name:
            errors.append({"row": i, "name": "—", "error": "Nombre vacío."})
            continue
        try:
            base_slug = slugify(name)
            slug      = base_slug
            counter   = 1
            while Brand.objects.filter(slug=slug).exclude(name__iexact=name).exists():
                slug = f"{base_slug}-{counter}"
                counter += 1

            brand, was_created = Brand.objects.update_or_create(
                name__iexact=name,
                defaults={
                    "name":      name,
                    "slug":      slug,
                    "website":   str(row.get("website", "")).strip(),
                    "is_active": parse_bool(row.get("is_active", True)),
                },
            )
            (created if was_created else updated).append(name)
        except Exception as e:
            errors.append({"row": i, "name": name, "error": str(e)})

    return {
        "created": created,
        "updated": updated,
        "errors":  errors,
        "total_processed": len(created) + len(updated),
    }


def import_brands_from_excel(file) -> dict:
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb.active
    headers_raw = [str(c.value or "").strip()
                   for c in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [normalize_header(h) for h in headers_raw]
    missing = [c for c in BRAND_REQUIRED if c not in headers]
    if missing:
        return {"error": f"Faltan columnas requeridas: {', '.join(missing)}"}
    mapped = [BRAND_COLUMN_MAP.get(h, h) for h in headers]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(mapped, row)))
    return _parse_brand_rows(rows)


def import_brands_from_csv(file) -> dict:
    content = file.read().decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(content))
    headers = [normalize_header(h) for h in (reader.fieldnames or [])]
    missing = [c for c in BRAND_REQUIRED if c not in headers]
    if missing:
        return {"error": f"Faltan columnas requeridas: {', '.join(missing)}"}
    rows = []
    for row in reader:
        mapped_row = {BRAND_COLUMN_MAP.get(normalize_header(k), normalize_header(k)): v
                      for k, v in row.items()}
        rows.append(mapped_row)
    return _parse_brand_rows(rows)


# ═══════════════════════════════════════════════════════════════
# COURSES
# ═══════════════════════════════════════════════════════════════

COURSE_COLUMN_MAP = {
    "titulo":              "title",
    "descripcion":         "description",
    "descripcion_corta":   "short_description",
    "precio":              "price",
    "nivel":               "level",
    "duracion_horas":      "duration_hours",
    "video_preview":       "preview_url",
    "preview_url":         "preview_url",
    "destacado":           "is_featured",
    "publicado":           "is_published",
    "publicada":           "is_published",
    "activo":              "is_published",
    "seo_titulo":          "seo_title",
    "seo_descripcion":     "seo_description",
}
COURSE_REQUIRED = ["titulo", "precio"]

LEVEL_MAP = {
    "principiante": "beginner",
    "beginner":     "beginner",
    "basico":       "beginner",
    "intermedio":   "intermediate",
    "intermediate": "intermediate",
    "avanzado":     "advanced",
    "advanced":     "advanced",
}


def _parse_course_rows(rows: list[dict]) -> dict:
    from apps.academy.models import Course  # import diferido

    created, updated, errors = [], [], []

    for i, row in enumerate(rows, start=2):
        title = str(row.get("title", "")).strip()
        if not title:
            errors.append({"row": i, "name": "—", "error": "Título vacío."})
            continue
        try:
            price = parse_decimal(row.get("price", 0))
            if price < 0:
                errors.append({"row": i, "name": title, "error": "Precio inválido."})
                continue

            base_slug = slugify(title)
            slug      = base_slug
            counter   = 1
            while Course.objects.filter(slug=slug).exclude(title__iexact=title).exists():
                slug = f"{base_slug}-{counter}"
                counter += 1

            level_raw = str(row.get("level", "beginner")).strip().lower()
            level     = LEVEL_MAP.get(level_raw, "beginner")

            defaults = {
                "title":              title,
                "slug":               slug,
                "description":        str(row.get("description", "")),
                "short_description":  str(row.get("short_description", ""))[:500],
                "price":              price,
                "level":              level,
                "preview_url":        str(row.get("preview_url", "")).strip(),
                "is_featured":        parse_bool(row.get("is_featured", False)),
                "is_published":       parse_bool(row.get("is_published", True)),
                "seo_title":          str(row.get("seo_title", ""))[:70],
                "seo_description":    str(row.get("seo_description", ""))[:160],
            }
            # duration_hours es opcional (puede no existir en Course)
            dh = row.get("duration_hours")
            if dh not in (None, "", "None"):
                try:
                    defaults["duration_hours"] = int(float(str(dh).replace(",", ".")))
                except (ValueError, TypeError):
                    pass

            # Aceptar sólo campos que existen en el modelo
            valid = {f.name for f in Course._meta.get_fields()}
            defaults = {k: v for k, v in defaults.items() if k in valid}

            course, was_created = Course.objects.update_or_create(
                title__iexact=title,
                defaults=defaults,
            )
            (created if was_created else updated).append(title)
        except Exception as e:
            errors.append({"row": i, "name": title, "error": str(e)})

    return {
        "created": created,
        "updated": updated,
        "errors":  errors,
        "total_processed": len(created) + len(updated),
    }


def import_courses_from_excel(file) -> dict:
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb.active
    headers_raw = [str(c.value or "").strip()
                   for c in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [normalize_header(h) for h in headers_raw]
    missing = [c for c in COURSE_REQUIRED if c not in headers]
    if missing:
        return {"error": f"Faltan columnas requeridas: {', '.join(missing)}"}
    mapped = [COURSE_COLUMN_MAP.get(h, h) for h in headers]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(mapped, row)))
    return _parse_course_rows(rows)


def import_courses_from_csv(file) -> dict:
    content = file.read().decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(content))
    headers = [normalize_header(h) for h in (reader.fieldnames or [])]
    missing = [c for c in COURSE_REQUIRED if c not in headers]
    if missing:
        return {"error": f"Faltan columnas requeridas: {', '.join(missing)}"}
    rows = []
    for row in reader:
        mapped_row = {COURSE_COLUMN_MAP.get(normalize_header(k), normalize_header(k)): v
                      for k, v in row.items()}
        rows.append(mapped_row)
    return _parse_course_rows(rows)
